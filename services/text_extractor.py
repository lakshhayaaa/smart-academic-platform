import io

import fitz
from docx import Document
from pptx import Presentation


def extract_pdf_text(file_bytes: bytes) -> str:
    document = fitz.open(
        stream=file_bytes,
        filetype="pdf"
    )

    text = ""

    for page in document:
        text += page.get_text()

    document.close()

    return text


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))

    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_ppt_text(file_bytes: bytes) -> str:
    presentation = Presentation(io.BytesIO(file_bytes))

    text = ""

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text